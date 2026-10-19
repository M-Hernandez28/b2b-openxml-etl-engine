"""
sprint_scheduler.py
-------------------
Automated Round-Robin Sprint Load Balancer.

Reads FileMetadata records (path, size_bytes, word_count, priority),
distributes them across N specialists using a stateful Round-Robin strategy,
and writes the schedule to Master_Tracker.xlsx with color-coded priority fills.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Literal

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Priority → background hex color
PRIORITY_FILLS = {
    "P1": PatternFill(start_color="FF4C4C", end_color="FF4C4C", fill_type="solid"),
    "P2": PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid"),
    "P3": PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid"),
}

HEADER_FILL = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)

COLUMNS = [
    "Specialist",
    "File",
    "Priority",
    "Size (KB)",
    "Word Count",
    "Segment Weight",
]


@dataclass
class FileMetadata:
    """Metadata for a single document to be scheduled."""
    path: Path
    size_bytes: int
    word_count: int
    priority: Literal["P1", "P2", "P3"] = "P3"

    @property
    def size_kb(self) -> float:
        return round(self.size_bytes / 1024, 2)

    @property
    def segment_weight(self) -> float:
        """Composite load score: word_count * priority multiplier."""
        multipliers = {"P1": 1.5, "P2": 1.2, "P3": 1.0}
        return round(self.word_count * multipliers.get(self.priority, 1.0), 1)


@dataclass
class ScheduleResult:
    """Summary of the scheduling run."""
    tracker_path: str
    total_files: int
    team_size: int
    specialist_summary: dict = field(default_factory=dict)

    def log_summary(self) -> None:
        logger.info(
            "Schedule written to %s | %d files across %d specialists",
            self.tracker_path, self.total_files, self.team_size,
        )
        for spec, stats in self.specialist_summary.items():
            logger.info("  %s → %d files | %d words", spec, stats["files"], stats["word_count"])


def _auto_column_widths(ws) -> None:
    """Set column widths based on maximum cell content length."""
    for col_idx, col_cells in enumerate(ws.iter_cols(), start=1):
        max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)


def schedule_sprint(
    files: List[FileMetadata],
    tracker_path: Path,
    team_size: int = 4,
    specialist_prefix: str = "Specialist",
) -> ScheduleResult:
    """
    Assign files to specialists using Round-Robin and write to Excel.

    Parameters
    ----------
    files : list of FileMetadata
        Documents to schedule. P1 items are sorted to the front.
    tracker_path : Path
        Output path for Master_Tracker.xlsx.
    team_size : int
        Number of specialists in the pod.
    specialist_prefix : str
        Label prefix for each specialist (e.g. 'Specialist' → 'Specialist_1').

    Returns
    -------
    ScheduleResult
    """
    # Sort P1 → P2 → P3 before assignment
    priority_order = {"P1": 0, "P2": 1, "P3": 2}
    sorted_files = sorted(files, key=lambda f: priority_order.get(f.priority, 3))

    specialists = [f"{specialist_prefix}_{i+1}" for i in range(team_size)]
    summary = {s: {"files": 0, "word_count": 0, "size_kb": 0.0} for s in specialists}
    rows = []

    for idx, doc in enumerate(sorted_files):
        specialist = specialists[idx % team_size]
        rows.append({
            "Specialist": specialist,
            "File": doc.path.name,
            "Priority": doc.priority,
            "Size (KB)": doc.size_kb,
            "Word Count": doc.word_count,
            "Segment Weight": doc.segment_weight,
        })
        summary[specialist]["files"] += 1
        summary[specialist]["word_count"] += doc.word_count
        summary[specialist]["size_kb"] += doc.size_kb

    # Build workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sprint_Schedule"

    # Header row
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows
    for row_idx, row in enumerate(rows, start=2):
        priority = row["Priority"]
        fill = PRIORITY_FILLS.get(priority)
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[col_name])
            if fill:
                cell.fill = fill
            cell.alignment = Alignment(horizontal="left")

    _auto_column_widths(ws)
    ws.freeze_panes = "A2"

    tracker_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(tracker_path)

    result = ScheduleResult(
        tracker_path=str(tracker_path),
        total_files=len(files),
        team_size=team_size,
        specialist_summary=summary,
    )
    result.log_summary()
    return result
